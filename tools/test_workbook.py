"""Self-check for workbook.py: python tools/test_workbook.py (uses a temp workspace, never yours)."""
import json, os, subprocess, sys, tempfile

TOOL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "workbook.py")
WS = tempfile.mkdtemp()
ENV = dict(os.environ, CAMPAIGN_WORKSPACE=WS)


def run(*args, ok=True):
    p = subprocess.run([sys.executable, TOOL, *args], env=ENV, capture_output=True, text=True)
    assert (p.returncode == 0) == ok, (args, p.stdout, p.stderr)
    return p.stdout + p.stderr


def write(spec, ok=True):
    path = os.path.join(WS, "spec.json")
    json.dump(spec, open(path, "w", encoding="utf-8"))
    return run("write", path, ok=ok)


def show(sheet):
    return [json.loads(l) for l in run("show", sheet).splitlines() if l.startswith("{")]


run("init", "jobs")
lead = {"Company": "Acme Re", "Role": "Actuarial Analyst", "Job URL": "https://www.acme.com/jobs/42?utm_source=x&req=42"}
write({"label": "t1", "add": {"Applications": [lead]}})
assert show("Applications")[0]["State"] == "Lead saved - not applied"

out = write({"add": {"Applications": [dict(lead, **{"Job URL": "https://acme.com/jobs/42/?req=42"})]}})
assert "same URL as JOB-0001" in out and len(show("Applications")) == 1, out

write({"add": {"Applications": [dict(lead, **{"Job URL": "https://acme.com/jobs/77"})]}})
assert len(show("Applications")) == 2 and "JOB-0002 may duplicate JOB-0001" in show("Review")[0]["Issue"]

assert "is an outcome" in write({"update": {"Applications": {"JOB-0001": {"State": "Applied"}}}}, ok=False)
write({"add": {"Event History": [{"Record ID": "JOB-0001", "Event type": "Applied", "Detail": "=HYPERLINK(\"x\")"}]},
       "update": {"Applications": {"JOB-0001": {"State": "Applied", "Applied date": "2026-11-02"}}}})
assert show("Applications")[0]["State"] == "Applied"
assert show("Event History")[0]["Detail"].startswith("=HYPERLINK")  # stored as text, not a formula

assert "must be one of" in write({"update": {"Applications": {"JOB-0002": {"State": "Sent"}}}}, ok=False)
assert "append-only" in write({"update": {"Event History": {"EV-0001": {"Detail": "x"}}}}, ok=False)
assert "YYYY-MM-DD" in write({"update": {"Applications": {"JOB-0002": {"Applied date": "2 Nov"}}}}, ok=False)

open(os.path.join(WS, ".write.lock"), "w").close()
assert "another writer" in write({"add": {"Applications": [{"Company": "B", "Role": "C"}]}}, ok=False)
os.remove(os.path.join(WS, ".write.lock"))
assert len(show("Applications")) == 2

run("init", "jobs,outreach,content,engagement,interview,resume")
assert show("Applications")[0]["State"] == "Applied"  # adding modules kept existing data
write({"add": {"Interviews": [{"Linked record": "JOB-0001", "Company": "Acme Re", "Round": "1", "Status": "Scheduled"}]}})
assert "is an outcome" in write({"update": {"Interviews": {"INT-0001": {"Status": "Done"}}}}, ok=False)
assert "unknown modules" in run("init", "astrology", ok=False)
assert len(os.listdir(os.path.join(WS, "backups"))) >= 5
# outcomes need evidence on new rows too, and events must point at real records
assert "cannot start as 'Applied'" in write({"add": {"Applications": [{"Company": "D", "Role": "E", "State": "Applied"}]}}, ok=False)
assert "does not exist" in write({"add": {"Event History": [{"Record ID": "JOB-0040", "Event type": "Applied"}]}}, ok=False)
assert "real date" in write({"update": {"Applications": {"JOB-0002": {"Applied date": "2026-13-45"}}}}, ok=False)

# hash-routed career sites: different jobs stay different
write({"add": {"Applications": [{"Company": "F", "Role": "G", "Job URL": "https://f.com/careers#/jobs/1"},
                                {"Company": "F", "Role": "H", "Job URL": "https://f.com/careers#/jobs/2"}]}})
assert sum(r["Company"] == "F" for r in show("Applications")) == 2

# a formatted blank block in Excel does not push new rows down
import openpyxl
live = os.path.join(WS, "campaign.xlsx")
wb = openpyxl.load_workbook(live)
for r in range(10, 40):
    wb["Applications"].cell(r, 2).font = openpyxl.styles.Font(bold=True)
wb.save(live)
write({"add": {"Applications": [{"Company": "I", "Role": "J"}]}})
ws = openpyxl.load_workbook(live)["Applications"]
assert [ws.cell(r, 2).value for r in range(2, 7)] == ["Acme Re", "Acme Re", "F", "F", "I"]

# refused writes leave no backup; Excel and LibreOffice locks stop writes
backups = len(os.listdir(os.path.join(WS, "backups")))
for name in ("~$campaign.xlsx", ".~lock.campaign.xlsx#"):
    open(os.path.join(WS, name), "w").close()
    assert "open in Excel or LibreOffice" in write({"add": {"Applications": [{"Company": "K", "Role": "L"}]}}, ok=False)
    os.remove(os.path.join(WS, name))
assert len(os.listdir(os.path.join(WS, "backups"))) == backups

# content openpyxl cannot read (a drawing someone added in Excel) blocks the write
import zipfile
with zipfile.ZipFile(live, "a") as z:
    z.writestr("xl/drawings/drawing1.xml", "<xdr:wsDr/>")
before = open(live, "rb").read()
assert "drop Excel content" in write({"add": {"Applications": [{"Company": "M", "Role": "N"}]}}, ok=False)
assert open(live, "rb").read() == before
print("all workbook checks passed")
