"""Tailored resumes that can only use facts from the owner's master resume.

  python tools/resume.py check  workspace/resume/variants/<label>.json
  python tools/resume.py render workspace/resume/variants/<label>.json   # -> .html and .pdf next to it

workspace/resume/master.json (built once from their real resume, every bullet confirmed):
  {"name": "", "contact": ["City", "email", "linkedin.com/in/..."],
   "experience": [{"id": "e1", "title": "", "org": "", "place": "", "dates": "",
                   "bullets": [{"id": "e1.1", "text": "", "source": "resume v3.pdf"}]}],
   "projects": [same shape as experience], "education": [{"school": "", "degree": "", "dates": ""}],
   "skills": [{"group": "", "items": ""}], "extras": [{"heading": "Awards", "lines": [""]}]}

variant.json (one per role family or per job):
  {"label": "", "for": "JOB-0004 or a role family", "page": "A4" or "Letter",
   "headline": "", "summary": "",
   "experience": [{"id": "e1", "bullets": ["e1.2", {"id": "e1.1", "text": "reworded for this role"}]}],
   "projects": [...same...], "skills": ["Group name", ...], "extras": ["Awards"], "education": true}

Rules the check enforces: every entry and bullet exists in master; a reworded bullet uses only numbers
in its master bullet; headline and summary use only numbers found somewhere in master; no Off limits
value from the workbook's Context sheet (one item per row) appears, unless the owner said yes to it for resumes and it is
listed in the variant's "allow": ["Acme Re"] (an employer name they hide in posts but not on a resume).
"""
import html, json, os, re, shutil, subprocess, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WS_DIR = os.environ.get("CAMPAIGN_WORKSPACE") or os.path.join(ROOT, "workspace")
NUM = re.compile(r"\d+(?:[.,]\d+)*")
BROWSERS = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"]


def nums(text):
    return {n.replace(",", "") for n in NUM.findall(text or "")}


def off_limits():
    path = os.path.join(WS_DIR, "campaign.xlsx")
    if not os.path.exists(path):
        return []
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True)["Context"]
    h = [c.value for c in next(ws.iter_rows(max_row=1))]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(h, r))
        if row.get("Category") == "Off limits" and row.get("Status") != "superseded" and row.get("Value"):
            out.append(str(row["Value"]).strip())  # one item per row, matched as written
    return out


def build(variant_path):
    """Resolve a variant against master. Returns (resume dict, problems)."""
    v = json.load(open(variant_path, encoding="utf-8"))
    m = json.load(open(os.path.join(WS_DIR, "resume", "master.json"), encoding="utf-8"))
    # ids (e1.2) and source notes (resume v3.pdf, user said 2026-10-31) are not facts
    problems, all_nums = [], nums(re.sub(r'"(id|source)": "[^"]*"', "", json.dumps(m, ensure_ascii=False)))
    out = {"name": m["name"], "contact": m.get("contact", []), "page": v.get("page", "A4"),
           "headline": v.get("headline", ""), "summary": v.get("summary", "")}
    for field in ("headline", "summary"):
        if extra := nums(v.get(field)) - all_nums:
            problems.append(f"{field} uses numbers not in master: {sorted(extra)}")
    for sec in ("experience", "projects"):
        entries = {e["id"]: e for e in m.get(sec, [])}
        out[sec] = []
        for pick in v.get(sec, []):
            e = entries.get(pick["id"])
            if not e:
                problems.append(f"{sec} {pick['id']} is not in master")
                continue
            bullets = {b["id"]: b["text"] for b in e.get("bullets", [])}
            texts = []
            for b in pick.get("bullets", []):
                bid, text = (b, None) if isinstance(b, str) else (b["id"], b["text"])
                if bid not in bullets:
                    problems.append(f"bullet {bid} is not in master")
                    continue
                if text and (extra := nums(text) - nums(bullets[bid])):
                    problems.append(f"bullet {bid} rewording adds numbers {sorted(extra)}; source says: {bullets[bid]}")
                texts.append(text or bullets[bid])
            out[sec].append({**{k: e.get(k, "") for k in ("title", "org", "place", "dates")}, "bullets": texts})
    groups = {g["group"]: g["items"] for g in m.get("skills", [])}
    out["skills"] = [(g, groups[g]) for g in v.get("skills", []) if g in groups]
    problems += [f"skills group {g!r} is not in master" for g in v.get("skills", []) if g not in groups]
    extras = {x["heading"]: x["lines"] for x in m.get("extras", [])}
    out["extras"] = [(x, extras[x]) for x in v.get("extras", []) if x in extras]
    out["education"] = m.get("education", []) if v.get("education", True) else []
    blob = json.dumps(out, ensure_ascii=False).lower()
    allowed = {a.lower() for a in v.get("allow", [])}  # owner said yes to these on a resume
    problems += [f"off-limits text appears: {w!r}" for w in off_limits() if w.lower() in blob and w.lower() not in allowed]
    return out, problems


def to_html(r):
    e = html.escape
    parts = [f"<h1>{e(r['name'])}</h1>"]
    if r["headline"]:
        parts.append(f"<p class=headline>{e(r['headline'])}</p>")
    parts.append(f"<p class=contact>{'  |  '.join(e(c) for c in r['contact'])}</p>")
    if r["summary"]:
        parts.append(f"<h2>Summary</h2><p>{e(r['summary'])}</p>")
    for sec, title in (("experience", "Experience"), ("projects", "Projects")):
        if r[sec]:
            parts.append(f"<h2>{title}</h2>")
            for x in r[sec]:
                where = ", ".join(e(s) for s in (x["org"], x["place"]) if s)
                parts.append(f"<div class=role><b>{e(x['title'])}</b>{' | ' + where if where else ''}<span>{e(x['dates'])}</span></div>")
                parts.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in x["bullets"]) + "</ul>")
    if r["skills"]:
        parts.append("<h2>Skills</h2>" + "".join(f"<p><b>{e(g)}:</b> {e(i)}</p>" for g, i in r["skills"]))
    if r["education"]:
        parts.append("<h2>Education</h2>" + "".join(
            f"<div class=role><b>{e(x.get('school', ''))}</b> | {e(x.get('degree', ''))}<span>{e(x.get('dates', ''))}</span></div>"
            for x in r["education"]))
    for heading, lines in r["extras"]:
        parts.append(f"<h2>{e(heading)}</h2>" + "".join(f"<p>{e(l)}</p>" for l in lines))
    css = (f"@page{{size:{r['page']};margin:14mm 15mm}}body{{font:10pt/1.35 'Segoe UI',Helvetica,Arial,sans-serif;color:#222;margin:0}}"
           "h1{font-size:18pt;margin:0;text-align:center}.headline{text-align:center;font-weight:600;color:#1f5f6b;margin:2px 0}"
           ".contact{text-align:center;color:#555;font-size:9pt;margin:2px 0 6px}h2{font-size:10.5pt;text-transform:uppercase;"
           "color:#1f5f6b;border-bottom:1px solid #bbb;margin:10px 0 4px;padding-bottom:1px}p{margin:2px 0}"
           ".role{display:flex;gap:4px;margin-top:5px}.role span{margin-left:auto;color:#555;white-space:nowrap}"
           "ul{margin:2px 0 0 16px;padding:0}li{margin:1px 0}")
    return f"<!doctype html><html lang=en><meta charset=utf-8><title>{e(r['name'])}</title><style>{css}</style><body>{''.join(parts)}</body></html>"


def render(variant_path):
    r, problems = build(variant_path)
    if problems:
        sys.exit("not rendered:\n- " + "\n- ".join(problems))
    base = os.path.splitext(os.path.abspath(variant_path))[0]
    open(base + ".html", "w", encoding="utf-8").write(to_html(r))
    if os.path.exists(base + ".pdf"):
        os.remove(base + ".pdf")  # a failed print must not leave the old PDF looking new
    exe = next((b for b in BROWSERS if os.path.exists(b)), None) or next(
        (shutil.which(n) for n in ("msedge", "google-chrome", "chromium", "chromium-browser", "chrome") if shutil.which(n)), None)
    if not exe:
        print(f"wrote {base}.html. No Chrome or Edge found: open it in a browser and Print > Save as PDF.")
        return
    subprocess.run([exe, "--headless", "--disable-gpu", "--no-pdf-header-footer", "--print-to-pdf-no-header",
                    f"--print-to-pdf={base}.pdf", "file:///" + (base + ".html").replace("\\", "/").lstrip("/")],
                   capture_output=True, timeout=120)
    print(f"wrote {base}.html and {base}.pdf" if os.path.exists(base + ".pdf") else
          f"wrote {base}.html; PDF failed: open it in a browser and Print > Save as PDF.")


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] not in ("check", "render"):
        sys.exit(__doc__)
    if sys.argv[1] == "check":
        _, p = build(sys.argv[2])
        if p:
            sys.exit("problems:\n- " + "\n- ".join(p))
        print("ok: every line traces to master.json")
    else:
        render(sys.argv[2])
