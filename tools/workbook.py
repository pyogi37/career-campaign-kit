"""The campaign workbook: one .xlsx in workspace/ that is the source of truth.

  python tools/workbook.py init jobs,outreach,content,engagement   # create, or add sheets for new modules
  python tools/workbook.py write spec.json                         # the only way agents change it
  python tools/workbook.py show Overview                           # counts at a glance
  python tools/workbook.py show Applications [--where "State=Applied"] [--limit 50]

spec.json = {
  "label": "daily-leads",
  "add":    {"Applications": [{"Company": "...", "Role": "...", "Job URL": "..."}],
             "Event History": [{"Record ID": "JOB-0004", "Event type": "Applied", "Detail": "..."}]},
  "update": {"Applications": {"JOB-0004": {"State": "Applied", "Applied date": "2026-11-02"}}}
}

Every write: exclusive lock, dated backup, save to a staging file, reopen and diff every cell,
check nobody changed the live file meanwhile, then replace it atomically.
"""
import datetime, hashlib, json, os, re, shutil, sys, uuid
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WS_DIR = os.environ.get("CAMPAIGN_WORKSPACE") or os.path.join(ROOT, "workspace")
LIVE = os.path.join(WS_DIR, "campaign.xlsx")
LOCK = os.path.join(WS_DIR, ".write.lock")

# sheet: (module, id prefix, status column, allowed statuses, columns after ID)
SHEETS = {
    "Applications": ("jobs", "JOB", "State",
        ["Lead saved - not applied", "Ready to apply", "Applied", "Interviewing", "Offer", "Accepted", "Declined", "Rejected", "Closed", "Skipped"],
        ["Company", "Role", "Job URL", "Location", "Source", "Captured", "State", "Fit (1-5)", "Fit notes and gaps",
         "Next action", "Resume used", "Applied date", "Last update", "Notes"]),
    "Event History": ("core", "EV", None, None,
        ["Date", "Record ID", "Company", "Event type", "Detail", "Evidence or link", "Source"]),
    "Outreach": ("outreach", "OUT", "Stage",
        ["Drafted", "Approved", "Invite sent", "Connected", "Message sent", "Replied", "Call booked", "Closed"],
        ["Person", "Company", "Profile URL", "Why them", "Linked record", "Stage", "Invite note", "First message",
         "Follow-up 1", "Follow-up 2", "Last sent", "Next follow-up", "Notes"]),
    "Content": ("content", "POST", "Status",
        ["Idea", "Planned", "Drafting", "Ready", "Posted", "Dropped"],
        ["Title", "Status", "Planned date", "Posted date", "URL", "Visual", "First line",
         "Impressions 7d", "Reactions 7d", "Comments 7d", "Profile views 7d", "Notes"]),
    "Engagement": ("engagement", "ENG", None, None,
        ["Date", "Person", "Company", "Type", "Post URL", "Note"]),
    "Context": ("core", "CTX", None, None,
        ["Category", "Fact", "Value", "Status", "Source", "Date recorded"]),
    "Review": ("core", "REV", None, None,
        ["Opened", "Priority", "Issue", "Unknown or conflict", "Next check", "Resolved"]),
}
DATE_COLS = {"Captured", "Applied date", "Last update", "Date", "Last sent", "Next follow-up", "Planned date",
             "Posted date", "Date recorded", "Opened"}
# A status change in these states is an outcome; it must arrive with an Event History row as evidence.
NEEDS_EVENT = {"Applied", "Interviewing", "Offer", "Accepted", "Declined", "Rejected",
               "Invite sent", "Message sent", "Replied", "Call booked", "Posted"}
TRACKING = re.compile(r"^(utm_.*|gh_src|ref|refid|trk|trackingid|source|src|lipi|fbclid|gclid)$", re.I)
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3A5F")


def today():
    return datetime.date.today().isoformat()


def canon_url(u):
    if not u:
        return ""
    p = urlsplit(u.strip())
    q = urlencode([(k, v) for k, v in parse_qsl(p.query) if not TRACKING.match(k)])
    return urlunsplit((p.scheme.lower(), p.netloc.lower().removeprefix("www."), p.path.rstrip("/"), q, ""))


def headers(ws):
    return [c.value for c in ws[1]]


def rows(ws):
    h = headers(ws)
    return [dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def make_sheet(wb, name):
    _, prefix, status_col, statuses, cols = SHEETS[name]
    ws = wb.create_sheet(name)
    ws.append(["ID"] + cols)
    for i, c in enumerate(ws[1], 1):
        c.font, c.fill = HEAD, FILL
        ws.column_dimensions[c.column_letter].width = 14 if i == 1 else 22
    ws.freeze_panes = "B2"
    if status_col:
        letter = ws.cell(1, cols.index(status_col) + 2).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(statuses) + '"', allow_blank=True)
        dv.add(f"{letter}2:{letter}5000")
        ws.add_data_validation(dv)


def refresh_overview(wb):
    ws = wb["Overview"] if "Overview" in wb.sheetnames else wb.create_sheet("Overview", 0)
    ws.delete_rows(1, ws.max_row)
    ws.append(["Campaign overview", f"refreshed {datetime.datetime.now():%Y-%m-%d %H:%M}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Counts are rewritten on every write. Edit the other sheets, not this one."])
    for name, (_, _, status_col, statuses, _) in SHEETS.items():
        if name not in wb.sheetnames:
            continue
        data = rows(wb[name])
        ws.append([])
        ws.append([name, len(data)])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        if status_col:
            for s in statuses:
                n = sum(1 for r in data if r.get(status_col) == s)
                if n:
                    ws.append([f"  {s}", n])
        if name == "Review":
            ws.append(["  open", sum(1 for r in data if not r.get("Resolved"))])
    ws.column_dimensions["A"].width = 32


def snapshot(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
            for ws in wb.worksheets if ws.title != "Overview"}


def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def transaction(label, change):
    """Run change(wb) -> (intended {(sheet, cell)}, report) under the shared write protocol."""
    os.makedirs(os.path.join(WS_DIR, "backups"), exist_ok=True)
    if os.path.exists(os.path.join(WS_DIR, "~$campaign.xlsx")):
        sys.exit("campaign.xlsx is open in Excel; close it and retry. Nothing written.")
    try:
        fd = os.open(LOCK, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError:
        sys.exit(f"another writer holds {LOCK}; nothing written. Delete it only if no agent is running.")
    run = uuid.uuid4().hex[:8]
    stage = os.path.join(WS_DIR, f"campaign.stage-{run}.xlsx")
    try:
        os.write(fd, f"{label} {run} {datetime.datetime.now().isoformat()}\n".encode())
        exists = os.path.exists(LIVE)
        base = sha(LIVE) if exists else None
        if exists:
            shutil.copy2(LIVE, os.path.join(WS_DIR, "backups", f"campaign_{today()}_{label}_{run}.xlsx"))
        before = snapshot(LIVE) if exists else {}
        wb = openpyxl.load_workbook(LIVE) if exists else openpyxl.Workbook()
        if not exists:
            wb.remove(wb.active)
        intended, report = change(wb)
        refresh_overview(wb)
        for ws in wb.worksheets:
            if ws.title != "Overview" and ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
        wb.save(stage)

        after = snapshot(stage)
        for sh in before:
            assert sh in after, f"sheet {sh} disappeared"
            diff = {(sh, k) for k in set(before[sh]) | set(after[sh]) if before[sh].get(k) != after[sh].get(k)}
            assert not diff - intended, f"unintended changes: {sorted(diff - intended)[:10]}"
        if exists and sha(LIVE) != base:
            sys.exit("campaign.xlsx changed while this write ran; nothing replaced. Re-read and retry.")
        os.replace(stage, LIVE)
        print(json.dumps(report, indent=1, ensure_ascii=False))
    finally:
        os.close(fd)
        os.remove(LOCK)
        if os.path.exists(stage):
            os.remove(stage)


def cmd_init(modules):
    wanted = set(modules) | {"core"}
    unknown = wanted - {m for m, *_ in SHEETS.values()}
    if unknown:
        sys.exit(f"unknown modules {unknown}; choose from jobs, outreach, content, engagement")

    def change(wb):
        added = [n for n, (m, *_) in SHEETS.items() if m in wanted and n not in wb.sheetnames]
        for n in added:
            make_sheet(wb, n)
        return set(), {"added sheets": added, "file": LIVE}  # new sheets only; existing ones must not change

    transaction("init", change)


def cmd_write(spec_path):
    spec = json.load(open(spec_path, encoding="utf-8"))
    label = re.sub(r"[^\w-]", "", spec.get("label", "write"))[:40] or "write"

    def change(wb):
        intended, report = set(), {"added": [], "updated": [], "skipped": [], "review": []}
        new_events = []

        def put(ws, row, col, value):
            h = headers(ws)
            if col not in h:
                raise SystemExit(f"{ws.title} has no column {col!r}; columns are {h}")
            if col in DATE_COLS and value and not re.fullmatch(r"\d{4}-\d{2}-\d{2}( \d{2}:\d{2})?", str(value)):
                raise SystemExit(f"{ws.title}.{col} must be YYYY-MM-DD, got {value!r}")
            _, _, status_col, statuses, _ = SHEETS.get(ws.title, (None,) * 5)
            if col == status_col and value not in statuses:
                raise SystemExit(f"{ws.title}.{col} must be one of {statuses}, got {value!r}")
            c = ws.cell(row, h.index(col) + 1, value)
            if isinstance(value, str) and value.startswith("="):
                c.data_type = "s"  # external text is never a formula
            intended.add((ws.title, c.coordinate))

        def next_id(ws, prefix):
            nums = [int(m.group(1)) for r in rows(ws) if (m := re.fullmatch(prefix + r"-(\d+)", str(r["ID"])))]
            return f"{prefix}-{max(nums, default=0) + 1:04d}"

        def add_row(name, values):
            ws, prefix = wb[name], SHEETS[name][1]
            rid = next_id(ws, prefix)
            r = ws.max_row + 1
            put(ws, r, "ID", rid)
            for k, v in values.items():
                put(ws, r, k, v)
            return rid

        for name, items in spec.get("add", {}).items():
            if name not in wb.sheetnames:
                raise SystemExit(f"no sheet {name!r}; enable its module with init first")
            for item in items:
                item = dict(item)
                item.pop("ID", None)
                if name == "Applications":
                    item.setdefault("State", "Lead saved - not applied")
                    item.setdefault("Captured", today())
                    existing = rows(wb[name])
                    url = canon_url(item.get("Job URL"))
                    if url and (hit := next((e for e in existing if canon_url(e.get("Job URL")) == url), None)):
                        report["skipped"].append(f"{item.get('Company')} / {item.get('Role')}: same URL as {hit['ID']}")
                        continue
                    same = [e["ID"] for e in existing if str(e.get("Company", "")).strip().lower() == str(item.get("Company", "")).strip().lower()
                            and str(e.get("Role", "")).strip().lower() == str(item.get("Role", "")).strip().lower()]
                elif name == "Event History":
                    item.setdefault("Date", today())
                    new_events.append(item.get("Record ID"))
                rid = add_row(name, item)
                report["added"].append(f"{name} {rid}")
                if name == "Applications" and same:
                    rev = add_row("Review", {"Opened": today(), "Priority": "P2", "Issue": f"{rid} may duplicate {', '.join(same)}",
                                             "Unknown or conflict": "Same company and role, different URL. Could be a new requisition.",
                                             "Next check": "Compare the two postings; close one or note why both stay."})
                    report["review"].append(rev)

        for name, changes in spec.get("update", {}).items():
            if name == "Event History":
                raise SystemExit("Event History is append-only; add a correcting event instead")
            ws = wb[name]
            ids = {str(ws.cell(r, 1).value): r for r in range(2, ws.max_row + 1)}
            status_col = SHEETS.get(name, (None,) * 5)[2]
            for rid, fields in changes.items():
                if rid not in ids:
                    raise SystemExit(f"{name} has no {rid}")
                if fields.get(status_col) in NEEDS_EVENT and rid not in new_events:
                    raise SystemExit(f"{rid} -> {fields[status_col]} is an outcome: add an Event History row with Record ID {rid} in the same spec")
                for k, v in fields.items():
                    if k == "ID":
                        raise SystemExit("IDs are stable; never change them")
                    put(ws, ids[rid], k, v)
                if name == "Applications" and "Last update" not in fields:
                    put(ws, ids[rid], "Last update", today())
                report["updated"].append(f"{name} {rid}")
        return intended, report

    transaction(label, change)


def cmd_show(name, where=None, limit=50):
    wb = openpyxl.load_workbook(LIVE, read_only=True)
    if name == "Overview":
        for r in wb[name].iter_rows(values_only=True):
            print("  ".join(str(v) for v in r if v is not None))
        return
    data = rows(wb[name])
    if where:
        k, v = where.split("=", 1)
        data = [r for r in data if str(r.get(k.strip())) == v.strip()]
    for r in data[-limit:]:
        print(json.dumps({k: v for k, v in r.items() if v is not None}, ensure_ascii=False, default=str))
    print(f"# {len(data)} rows in {name}" + (f" where {where}" if where else ""), file=sys.stderr)


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a or a[0] not in ("init", "write", "show"):
        sys.exit(__doc__)
    if a[0] == "init":
        cmd_init(a[1].split(",") if len(a) > 1 else ["jobs"])
    elif a[0] == "write":
        cmd_write(a[1])
    else:
        opt = dict(zip(a[2::2], a[3::2]))
        cmd_show(a[1], opt.get("--where"), int(opt.get("--limit", 50)))
